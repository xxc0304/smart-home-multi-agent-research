from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "智慧家庭Agent综述_新稿_v0.1.md"
OUTPUT = ROOT / "智慧家庭Agent综述_教师审阅版_v0.2.docx"


def set_run_font(run, east_asia="宋体", size=10.5, bold=None, italic=None, color=None):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color="666666")


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "宋体")
    r_pr.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "19")
    r_pr.append(size)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def inline_parts(text):
    pattern = re.compile(r"(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            yield ("text", text[pos:match.start()])
        token = match.group(0)
        if token.startswith("["):
            m = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", token)
            yield ("link", m.group(1), m.group(2))
        elif token.startswith("**"):
            yield ("bold", token[2:-2])
        elif token.startswith("*"):
            yield ("italic", token[1:-1])
        else:
            yield ("text", token[1:-1])
        pos = match.end()
    if pos < len(text):
        yield ("text", text[pos:])


def add_rich_text(paragraph, text, size=10.5):
    for part in inline_parts(text):
        if part[0] == "link":
            add_hyperlink(paragraph, part[1], part[2])
        else:
            run = paragraph.add_run(part[1])
            set_run_font(run, size=size, bold=(part[0] == "bold"), italic=(part[0] == "italic"))


def clean_markdown_text(text, citation_map):
    def grouped_replace(match):
        keys = re.findall(r"H(\d{2})", match.group(0))
        numbers = ", ".join(citation_map.get(f"H{key}", key) for key in keys)
        return f"[{numbers}]"

    text = re.sub(r"\[(?:H\d{2})(?:,\s*H\d{2})+\]", grouped_replace, text)

    def bracket_replace(match):
        key = f"H{match.group(1)}"
        return f"[{citation_map.get(key, match.group(1))}]"

    text = re.sub(r"\[H(\d{2})\]", bracket_replace, text)

    def plain_replace(match):
        key = f"H{match.group(1)}"
        return f"[{citation_map.get(key, key)}]"

    text = re.sub(r"(?<![A-Za-z])H(\d{2})(?![A-Za-z0-9])", plain_replace, text)
    text = text.replace("[Hxx]", "[内部编号]")
    return text


def split_table_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def add_table(doc, rows, citation_map):
    header = [clean_markdown_text(c, citation_map) for c in split_table_row(rows[0])]
    body_rows = []
    for line in rows[2:]:
        cells = split_table_row(line)
        if cells:
            body_rows.append([clean_markdown_text(c, citation_map) for c in cells])
    cols = len(header)
    table = doc.add_table(rows=1, cols=cols)
    table.style = "Table Grid"
    table.autofit = False
    widths = {
        "类型": [0.18, 0.28, 0.30, 0.24],
        "组织形式": [0.18, 0.28, 0.30, 0.24],
        "应用任务": [0.19, 0.28, 0.31, 0.22],
        "评价维度": [0.19, 0.34, 0.47],
        "评测阶段": [0.16, 0.17, 0.18, 0.20, 0.29],
        "演化阶段": [0.15, 0.14, 0.15, 0.18, 0.20, 0.18],
    }.get(header[0], [1.0 / cols] * cols)
    available = 16.0

    def fill_row(row, values, is_header=False, row_index=0):
        for idx, cell in enumerate(row.cells):
            cell.width = Cm(available * widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if not is_header and row_index % 2 == 1:
                set_cell_shading(cell, "F5F8FB")
            if is_header:
                set_cell_shading(cell, "D9EAF7")
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            for part in inline_parts(values[idx] if idx < len(values) else ""):
                if part[0] == "link":
                    add_hyperlink(p, part[1], part[2])
                else:
                    run = p.add_run(part[1])
                    set_run_font(run, size=8.2, bold=is_header or part[0] == "bold", italic=part[0] == "italic")

    fill_row(table.rows[0], header, is_header=True)
    set_repeat_table_header(table.rows[0])
    for idx, values in enumerate(body_rows):
        row = table.add_row()
        fill_row(row, values, row_index=idx)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    ref_ids = []
    for line in lines:
        m = re.match(r"^- \[(H\d{2})\]", line)
        if m and m.group(1) not in ref_ids:
            ref_ids.append(m.group(1))
    citation_map = {key: str(idx + 1) for idx, key in enumerate(ref_ids)}

    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(24)
    section.bottom_margin = Mm(22)
    section.left_margin = Mm(25)
    section.right_margin = Mm(25)
    add_page_number(section.footer.paragraphs[0])

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)
    for name, size, before, after in (("Heading 1", 14, 12, 6), ("Heading 2", 12, 9, 4), ("Heading 3", 11, 6, 3)):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = None
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title_p = doc.add_paragraph(style="Title")
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("智慧家庭智能体系统研究综述")
    set_run_font(title_run, east_asia="黑体", size=18, bold=True)
    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle_p.add_run("从家庭整体 Agent 到大语言模型驱动的任务执行与多智能体协同")
    set_run_font(subtitle_run, east_asia="黑体", size=12, bold=False)
    subtitle_p.paragraph_format.space_after = Pt(14)

    idx = 0
    pending_table = []
    while idx < len(lines):
        line = lines[idx]
        if line.startswith(">"):
            idx += 1
            continue
        if line.startswith("# "):
            idx += 1
            continue
        if line.startswith("|"):
            pending_table = [line]
            idx += 1
            while idx < len(lines) and lines[idx].startswith("|"):
                pending_table.append(lines[idx])
                idx += 1
            if len(pending_table) >= 2:
                add_table(doc, pending_table, citation_map)
            continue
        if not line.strip():
            idx += 1
            continue
        m = re.match(r"^(#{2,4})\s+(.+)$", line)
        if m:
            level = len(m.group(1)) - 1
            text = clean_markdown_text(m.group(2), citation_map)
            p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
            add_rich_text(p, text, size={1: 14, 2: 12, 3: 11}[min(level, 3)])
            idx += 1
            continue
        m = re.match(r"^(\d+)\.\s+(.+)$", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Cm(0.6)
            add_rich_text(p, clean_markdown_text(m.group(2), citation_map))
            idx += 1
            continue
        m = re.match(r"^[-*]\s+(.+)$", line)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.6)
            add_rich_text(p, clean_markdown_text(m.group(1), citation_map))
            idx += 1
            continue
        text = clean_markdown_text(line, citation_map)
        p = doc.add_paragraph()
        if text.startswith("**关键词：**"):
            run = p.add_run("关键词：")
            set_run_font(run, bold=True)
            add_rich_text(p, text.replace("**关键词：**", "", 1).strip())
        else:
            add_rich_text(p, text)
        idx += 1

    doc.core_properties.title = "智慧家庭智能体系统研究综述"
    doc.core_properties.subject = "智慧家庭 Agent 研究综述教师审阅稿"
    doc.core_properties.author = ""
    doc.core_properties.comments = ""
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
