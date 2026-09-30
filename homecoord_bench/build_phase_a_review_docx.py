"""Build the standalone Word form for independent C1/C3 phase-A review."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "C1_C3_PHASE_A_BLIND_REVIEW_2026-09-28.docx"

NAVY = "203864"
BLUE = "D9EAF7"
PALE = "F2F6FA"
GRAY = "5B6573"
GRID = "B8C4D1"
FONT = "Microsoft YaHei"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_margins(cell, top=80, start=95, bottom=80, end=95) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_cell_text(cell, text: str, *, bold=False, color=None, size=8.8) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text)
    run.bold = bold
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def style_paragraph(p, *, size=10, color="222222", bold=False, after=5, before=0):
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.1
    for run in p.runs:
        run.font.name = FONT
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
    return p


def add_para(doc, text="", *, size=10, color="222222", bold=False, after=5, before=0):
    p = doc.add_paragraph()
    p.add_run(text)
    return style_paragraph(p, size=size, color=color, bold=bold, after=after, before=before)


def add_heading(doc, text: str, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    if level == 1:
        size, color = 15, NAVY
    else:
        size, color = 11.5, NAVY
    style_paragraph(p, size=size, color=color, bold=True, after=5, before=4)
    return p


def add_bullet(doc, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    style_paragraph(p, size=9.5, after=3)
    return p


def add_prompt_block(doc, prompt: str, lines=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    p.add_run(prompt)
    style_paragraph(p, size=9.5, bold=True, after=2)
    for _ in range(lines):
        blank = doc.add_paragraph("________________________________________________________________________________")
        style_paragraph(blank, size=8, color="A8B1BC", after=4)


def add_grid_table(doc, headers, rows, widths=None, font_size=8.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    for i, heading in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], heading, bold=True, color="FFFFFF", size=font_size)
        set_cell_shading(table.rows[0].cells[i], NAVY)
    set_repeat_table_header(table.rows[0])
    for values in rows:
        row = table.add_row()
        set_no_split(row)
        for i, value in enumerate(values):
            set_cell_text(row.cells[i], value, size=font_size)
            if i == 0:
                set_cell_shading(row.cells[i], PALE)
        if widths:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    if widths:
        for cell, width in zip(table.rows[0].cells, widths):
            cell.width = Inches(width)
    return table


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    for name in ("Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("HomeCoord-Bench C1/C3 第一阶段独立审核表")
    style_paragraph(title, size=19, color=NAVY, bold=True, after=4)
    subtitle = add_para(doc, "任务语义、冲突定义与参数可信度", size=10.5,
                        color=GRAY, after=10)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_para(doc, "审核者代号：________________    日期：________________    专业方向：________________________________", size=9.5, after=8)
    add_heading(doc, "审核目标与独立性要求", 1)
    add_para(doc, "请只判断任务草稿是否自然、可执行、可评分，以及其中的设备和资源条件是否有依据。当前阶段不评判协调策略表现。两位审核者分别填写、提交并锁定意见前，不相互讨论。", size=9.5)
    add_para(doc, "请只使用随附的五条 JSON 任务草稿和本表。提交第一阶段意见前，不查看项目仓库里的实验结果、作者的参数来源判断、策略报告或其他审核者意见。当前数值若缺少实物或现场依据，可以建议保留为合成任务、改成抽象资源任务、修订或剔除。", size=9.5)

    add_heading(doc, "任务范围", 1)
    add_heading(doc, "C1  同一 HVAC 设备的动作时序", 2)
    add_para(doc, "一条草稿中，舒适任务请求空调制冷，能源任务请求关闭同一设备；配对草稿将能源动作安排在制冷服务之后。请审核任务目标、互斥关系、必要服务以及对照是否只改变关键时序条件。", size=9.2)
    add_heading(doc, "C3  照明与 HVAC 共用容量", 2)
    add_para(doc, "三条草稿包含书房灯 0.5 kW、卧室 HVAC 1.0 kW，并将共享容量设为低于、等于或高于两者合计的 1.5 kW。请审核容量口径和数值自然性。草稿没有附对应设备铭牌或家庭回路实测记录；若不能支持真实设备解释，可建议改作抽象共享资源任务。", size=9.2)

    add_heading(doc, "填写方式", 1)
    for item in [
        "在每项下选择：通过、需修订、剔除或无法判定。",
        "说明判断依据，尽可能指明 JSON 字段、设备规格、专业适用范围或反例；证据不足时写明“缺证”。",
        "指出可能改变实验解释的缺陷，并给出最小可行修订建议。",
        "确认所审版本与本表末尾 SHA-256 文件指纹一致。",
    ]:
        add_bullet(doc, item)
    add_para(doc, "第一阶段结束后，研究者会收回两份独立意见并锁定记录，再另行提供实验结果供第二阶段评估。", size=9, color=GRAY, after=0)

    doc.add_page_break()
    add_heading(doc, "C1 审核  HVAC 命令冲突", 1)
    add_para(doc, "请根据随附的 HC-PAIR-C1-HVAC-CONFLICT 与 HC-PAIR-C1-HVAC-SAFE 两条草稿填写。", size=9.5)
    add_prompt_block(doc, "1. 用户任务、舒适目标和能源任务的必要服务是否清楚、自然？", 2)
    add_prompt_block(doc, "2. 两个动作是否确实作用于同一设备且互斥？配对是否只改变请求时序？", 2)
    add_prompt_block(doc, "3. 工具、参数、权限、动作持续时间、回执和完成效果是否足以实际执行及评分？", 2)
    add_prompt_block(doc, "4. “降低峰值能耗”是否有明确的峰值时段和成功判据？若没有，应如何修改或降级？", 2)
    add_prompt_block(doc, "5. 草稿中的温度在数秒内从 29°C 降到 24°C。该效果是否可信？若不可信，任务可否只用于 HVAC 命令冲突诊断？", 3)
    add_prompt_block(doc, "C1 阶段判断与必须修改的 JSON 字段或措辞：", 3)

    doc.add_page_break()
    add_heading(doc, "C3 审核  共享容量任务", 1)
    add_para(doc, "请根据随附的 HC-PAIR-C3-HOME-POWER 三条草稿填写。", size=9.5)
    add_prompt_block(doc, "1. 书房照明与卧室制冷作为同时发生的用户任务是否自然？每个 Agent 的必要服务是否明确？", 2)
    add_prompt_block(doc, "2. 0.5 kW 灯具和 1.0 kW HVAC 的功率值适合描述普通家居设备吗？请给出依据或指出缺证。", 2)
    add_prompt_block(doc, "3. 1.2、1.5、1.7 kW 表示家庭总功率、单一回路上限，还是抽象共享资源容量？当前定义足够清楚吗？", 2)
    add_prompt_block(doc, "4. 三条配对是否只改变共享容量？两个单项负载各自可行，联合负载在低／临界／高容量下是否符合预期？", 2)
    add_prompt_block(doc, "5. Agent 与协调规则应分别看到哪些容量及任务功率信息？如何避免信息优势不公平？", 2)
    add_prompt_block(doc, "C3 阶段判断与必须修改的 JSON 字段或措辞：", 3)

    doc.add_page_break()
    add_heading(doc, "整体判断与提交", 1)
    add_para(doc, "请分别判断下列评测设计要求是否满足，并写明理由。", size=9.5)
    rubric = [
        ("任务目标和必要动作无歧义", "[ ] 通过  [ ] 修订  [ ] 不适用", ""),
        ("冲突定义有操作性，配对只改变一个关键条件", "[ ] 通过  [ ] 修订  [ ] 不适用", ""),
        ("单项与联合任务的物理可行性清楚", "[ ] 通过  [ ] 修订  [ ] 不适用", ""),
        ("设备数值有依据，或明确标成合成／抽象条件", "[ ] 通过  [ ] 修订  [ ] 不适用", ""),
        ("Agent、协调规则可见的信息公平且写明", "[ ] 通过  [ ] 修订  [ ] 不适用", ""),
        ("评分区分安全、必要服务、最终目标与时延", "[ ] 通过  [ ] 修订  [ ] 不适用", ""),
    ]
    add_grid_table(doc, ["审核项", "判定", "依据／建议"], rubric,
                   widths=(2.45, 1.55, 2.83), font_size=8.0)
    add_heading(doc, "最终建议", 2)
    add_para(doc, "C1： [ ] 保留为命令冲突诊断  [ ] 修订后再审  [ ] 剔除  [ ] 无法判断", size=9.3)
    add_para(doc, "C3： [ ] 保留为抽象容量任务  [ ] 有来源后作物理任务  [ ] 修订后再审  [ ] 剔除  [ ] 无法判断", size=9.3)
    add_prompt_block(doc, "最强反例，或你认为当前材料不足以冻结的主要理由：", 2)
    add_para(doc, "审核者签名／代号：________________________    日期：________________________", size=9.3, after=8)

    add_heading(doc, "草稿版本核对", 2)
    hashes = [
        ("C1-HVAC-CONFLICT", "BEF8CC8403EB58C17BD2498008680F58F522059641B1899D47ED8DDE61F25088"),
        ("C1-HVAC-SAFE", "62AC9992444C4D47C60F46E6A50947CC821693293F56CE767380A99F77C9CC97"),
        ("C3-OVER-CAPACITY", "FF74DFF8649332A058CD3C0BA80B69E4569CD34F7BEF7EEA7522882AC49BAC3A"),
        ("C3-EXACT-CAPACITY", "F2574F9A79E9892FD6E426CF4B3248628DAA525CFB2177A6958A3CBAB9B86054"),
        ("C3-SAFE-PARALLEL", "C333E19246D3FB474AD310CDA852E875FCAB2C6189B73780D523CB6473C2950B"),
    ]
    ht = doc.add_table(rows=1, cols=2)
    ht.style = "Table Grid"
    ht.autofit = False
    for i, text in enumerate(("文件简称", "SHA-256")):
        set_cell_text(ht.rows[0].cells[i], text, bold=True, color="FFFFFF", size=7.2)
        set_cell_shading(ht.rows[0].cells[i], NAVY)
    for short, digest in hashes:
        row = ht.add_row()
        set_no_split(row)
        set_cell_text(row.cells[0], short, size=7.0)
        set_cell_text(row.cells[1], digest, size=6.5)
    for row in ht.rows:
        row.cells[0].width = Inches(1.65)
        row.cells[1].width = Inches(5.18)

    doc.core_properties.title = "HomeCoord-Bench C1 C3 第一阶段独立审核表"
    doc.core_properties.subject = "盲审 C1 HVAC 命令冲突与 C3 共享容量任务草稿"
    doc.core_properties.author = "HomeCoord-Bench research team"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
